"""
Kokborok Vocabulary Construction & Corpus Processing Script.

Phase 10: Grounded in Phase 9 Empirical Findings:
1. Extracts all 4,020 valid sentences from the 6 read-only Kokborok Excel files.
2. Applies canonical Kokborok normalization (NFC, mojibake repair, LID tag stripping, whitespace).
3. Saves normalized sentences to data/processed/kokborok/sentences.txt.
4. Tokenizes sentences into words and characters using KokborokTokenizer.
5. Builds standalone Kokborok word vocabulary (data/processed/kokborok/word_vocab.json)
   and character vocabulary (data/processed/kokborok/char_vocab.json).
6. Preserves isolation: does NOT merge into the three-language vocabulary (reserved for Phase 12).
"""

from __future__ import annotations

import json
import os
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from preprocessing.kokborok_normalize import normalize_kokborok_text
from preprocessing.kokborok_tokenizer import KokborokTokenizer

SPECIAL_TOKENS = ["<PAD>", "<UNK>", "<BOS>", "<EOS>"]


def build_vocab_dict(
    counts: Counter[str],
    min_freq: int = 1,
    max_size: int | None = None,
    special_tokens: list[str] = SPECIAL_TOKENS,
) -> dict[str, Any]:
    """
    Build vocabulary dictionary with token_to_id, id_to_token, and frequencies.
    """
    token_to_id: dict[str, int] = {}
    id_to_token: dict[int, str] = {}

    for idx, sp_tok in enumerate(special_tokens):
        token_to_id[sp_tok] = idx
        id_to_token[idx] = sp_tok

    current_id = len(special_tokens)
    # Sort by frequency descending, then alphabetically for determinism
    sorted_tokens = sorted(counts.items(), key=lambda x: (-x[1], x[0]))

    for token, freq in sorted_tokens:
        if freq < min_freq:
            continue
        if max_size and current_id >= max_size:
            break
        token_to_id[token] = current_id
        id_to_token[current_id] = token
        current_id += 1

    return {
        "special_tokens": special_tokens,
        "vocab_size": len(token_to_id),
        "min_freq": min_freq,
        "token_to_id": token_to_id,
        "id_to_token": id_to_token,
        "frequencies": {tok: counts[tok] for tok, _ in sorted_tokens if counts[tok] >= min_freq},
    }


def extract_all_kokborok_sentences(dataset_dir: Path) -> list[dict[str, str]]:
    """
    Extract all raw sentences across the 6 files with file provenance.
    """
    records: list[dict[str, str]] = []

    # 1. English-Kokborok 3rd Dataset.xlsx
    f1 = dataset_dir / "English-Kokborok 3rd Dataset.xlsx"
    wb1 = openpyxl.load_workbook(f1, data_only=True)
    s1 = wb1["English-Kokborok 3rd Dataset"]
    for r in list(s1.iter_rows(values_only=True))[1:]:
        if len(r) > 1 and r[1] and str(r[1]).strip():
            records.append({
                "source_file": f1.name,
                "sheet": "English-Kokborok 3rd Dataset",
                "category": "parallel_translation",
                "raw_text": str(r[1]).strip(),
            })

    # 2. Kokborok second dataset (1).xlsx
    f2 = dataset_dir / "Kokborok second dataset (1).xlsx"
    wb2 = openpyxl.load_workbook(f2, data_only=True)
    s2 = wb2["Sheet1"]
    for r in list(s2.iter_rows(values_only=True))[1:]:
        if len(r) > 1 and r[1] and str(r[1]).strip():
            records.append({
                "source_file": f2.name,
                "sheet": "Sheet1",
                "category": "parallel_translation",
                "raw_text": str(r[1]).strip(),
            })

    # 3. Kokborok Sentiment data.xlsx
    f3 = dataset_dir / "Kokborok Sentiment data.xlsx"
    wb3 = openpyxl.load_workbook(f3, data_only=True)
    for sname in ["Sentiment Dataset", "Sarcasm Dataset"]:
        s3 = wb3[sname]
        for r in list(s3.iter_rows(values_only=True))[1:]:
            if len(r) > 0 and r[0] and str(r[0]).strip():
                records.append({
                    "source_file": f3.name,
                    "sheet": sname,
                    "category": "sentiment_sarcasm",
                    "raw_text": str(r[0]).strip(),
                })

    # 4. Kokborok Translations (1).xlsx
    f4 = dataset_dir / "Kokborok Translations (1).xlsx"
    wb4 = openpyxl.load_workbook(f4, data_only=True)
    s4 = wb4["Sheet1"]
    for r in list(s4.iter_rows(values_only=True))[1:]:
        if len(r) > 1 and r[1] and str(r[1]).strip():
            records.append({
                "source_file": f4.name,
                "sheet": "Sheet1",
                "category": "parallel_translation",
                "raw_text": str(r[1]).strip(),
            })

    # 5. Kokborok-dataset.xlsx
    f5 = dataset_dir / "Kokborok-dataset.xlsx"
    wb5 = openpyxl.load_workbook(f5, data_only=True)
    s5 = wb5["Sheet1"]
    for r in s5.iter_rows(values_only=True):
        if len(r) > 0 and r[0] and str(r[0]).strip():
            records.append({
                "source_file": f5.name,
                "sheet": "Sheet1",
                "category": "code_switched_lid",
                "raw_text": str(r[0]).strip(),
            })

    # 6. Question Answering.xlsx
    # In Phase 9: Confirmed unpopulated (0 answers)

    return records


def run_pipeline() -> dict[str, Any]:
    dataset_dir = PROJECT_ROOT / "Kokborok_datasets"
    output_dir = PROJECT_ROOT / "data" / "processed" / "kokborok"
    output_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 80)
    print("PHASE 10: KOKBOROK PIPELINE ADAPTATION & VOCABULARY CONSTRUCTION")
    print("=" * 80)

    # 1. Extraction
    print("\n[Step 1] Extracting raw sentences from Kokborok_datasets...")
    raw_records = extract_all_kokborok_sentences(dataset_dir)
    print(f"  Total records extracted: {len(raw_records)}")

    # 2. Normalization & Tokenization
    print("\n[Step 2] Normalizing text and tokenizing words and characters...")
    tok = KokborokTokenizer()

    normalized_sentences: list[str] = []
    word_counts: Counter[str] = Counter()
    char_counts: Counter[str] = Counter()
    total_tokens = 0

    for rec in raw_records:
        strip_lid = rec["category"] == "code_switched_lid"
        norm_text = normalize_kokborok_text(rec["raw_text"], strip_lid_tags=strip_lid)
        if not norm_text:
            continue

        normalized_sentences.append(norm_text)
        words = tok.tokenize_words(norm_text)
        total_tokens += len(words)

        for w in words:
            word_counts[w] += 1
            for ch in tok.tokenize_characters(w):
                char_counts[ch] += 1

    print(f"  Valid normalized sentences: {len(normalized_sentences)}")
    print(f"  Total word tokens: {total_tokens:,d}")
    print(f"  Unique word types (raw): {len(word_counts):,d}")
    print(f"  Unique characters: {len(char_counts):,d}")

    # 3. Write sentences.txt
    sentences_file = output_dir / "sentences.txt"
    with open(sentences_file, "w", encoding="utf-8") as f:
        for s in normalized_sentences:
            f.write(s + "\n")
    print(f"\n[Step 3] Normalized sentences written to: {sentences_file.relative_to(PROJECT_ROOT)}")

    # 4. Build and save word_vocab.json
    print("\n[Step 4] Building word vocabulary (min_freq=1 to preserve rare entities)...")
    word_vocab = build_vocab_dict(word_counts, min_freq=1, max_size=30000)
    word_vocab_file = output_dir / "word_vocab.json"
    with open(word_vocab_file, "w", encoding="utf-8") as f:
        json.dump(word_vocab, f, indent=2, ensure_ascii=False)
    print(f"  Word vocab size (including 4 special tokens): {word_vocab['vocab_size']:,d}")
    print(f"  Word vocab saved to: {word_vocab_file.relative_to(PROJECT_ROOT)}")

    # 5. Build and save char_vocab.json
    print("\n[Step 5] Building character vocabulary (min_freq=1)...")
    char_vocab = build_vocab_dict(char_counts, min_freq=1)
    char_vocab_file = output_dir / "char_vocab.json"
    with open(char_vocab_file, "w", encoding="utf-8") as f:
        json.dump(char_vocab, f, indent=2, ensure_ascii=False)
    print(f"  Char vocab size (including 4 special tokens): {char_vocab['vocab_size']:,d}")
    print(f"  Char vocab saved to: {char_vocab_file.relative_to(PROJECT_ROOT)}")

    results = {
        "total_extracted_records": len(raw_records),
        "normalized_sentences": len(normalized_sentences),
        "total_word_tokens": total_tokens,
        "word_vocab_size": word_vocab["vocab_size"],
        "char_vocab_size": char_vocab["vocab_size"],
        "word_vocab_path": str(word_vocab_file.relative_to(PROJECT_ROOT)),
        "char_vocab_path": str(char_vocab_file.relative_to(PROJECT_ROOT)),
        "sentences_path": str(sentences_file.relative_to(PROJECT_ROOT)),
    }
    return results


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    run_pipeline()
